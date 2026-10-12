from pathlib import Path
import json, re, math
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent
TMP=ROOT.parent.parent/'tmp/dp-handbook'
TMP.mkdir(parents=True,exist_ok=True)
NAMES=['01_动态规划_从零建立模型','02_动态规划_提高组建模与优化','03_动态规划_独立练习与验收','04_动态规划_妈妈陪读与教练核验']
FONT='Noto Serif CJK SC'; HEAD_FONT='Noto Sans CJK SC'; LATIN='Calibri'; MONO='Consolas'
PAGE_MAP=json.loads((TMP/'page_map.json').read_text()) if (TMP/'page_map.json').exists() else {}

def set_font(style,name,size,bold=False,east_asia=FONT):
 style.font.name=name;style.font.size=Pt(size);style.font.bold=bold;style.font.italic=False;style.font.color.rgb=RGBColor(0,0,0)
 rf=style.element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),east_asia)
 for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  rf.attrib.pop(qn('w:'+a),None)

def add_bookmark(p,name,ident):
 a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(ident));a.set(qn('w:name'),name)
 b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(ident));p._p.insert(0,a);p._p.append(b)

def link(p,label,anchor):
 h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),anchor)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'000000');pr.append(col)
 r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def add_field(p,instr):
 r=p.add_run();c=OxmlElement('w:fldSimple');c.set(qn('w:instr'),instr);r._r.addnext(c)

def diagram(kind):
 path=TMP/(kind+'.png')
 im=Image.new('RGB',(1500,430),'white');d=ImageDraw.Draw(im)
 font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Songti.ttc',31)
 fsmall=ImageFont.truetype('/System/Library/Fonts/Supplemental/Songti.ttc',26)
 def node(x,y,label):
  d.rounded_rectangle((x-94,y-40,x+94,y+40),radius=14,outline='#525252',width=3,fill='#f6f6f6')
  d.text((x,y),label,font=font,anchor='mm',fill='black')
 def edge(a,b,arrow=False):
  dx=b[0]-a[0];dy=b[1]-a[1];t=1/max(abs(dx)/94,abs(dy)/40)
  x,y=a[0]+dx*t,a[1]+dy*t;xx,yy=b[0]-dx*t,b[1]-dy*t
  d.line((x,y,xx,yy),fill='#555555',width=3)
  if arrow:
   t=math.atan2(yy-y,xx-x)
   d.polygon([(xx,yy),(xx-15*math.cos(t-.45),yy-15*math.sin(t-.45)),(xx-15*math.cos(t+.45),yy-15*math.sin(t+.45))],fill='#555555')
 if kind=='recursion':
  for a,b in [((760,60),(455,180)),((760,60),(1050,180)),((455,180),(275,335)),((455,180),(640,335)),((1050,180),(945,335)),((1050,180),(1280,335))]:edge(a,b,True)
  for x,y,t in [(760,60,'求第5级'),(455,180,'求第4级'),(1050,180,'求第3级'),(275,335,'求第3级'),(640,335,'求第2级'),(945,335,'求第2级'),(1280,335,'求第1级')]:node(x,y,t)
  d.text((745,408),'重复出现的相同问题 可以共享已经算出的答案',font=fsmall,anchor='mm',fill='black')
 elif kind=='dag':
  for a,b in [((300,215),(750,70)),((300,215),(750,360)),((750,70),(1200,215)),((750,360),(1200,215))]:edge(a,b,True)
  for x,y,t in [(300,215,'点1'),(750,70,'点2'),(750,360,'点3'),(1200,215,'点4')]:node(x,y,t)
  d.text((750,215),'两条路径分别经点2与点3',font=fsmall,anchor='mm',fill='black')
 elif kind=='tree':
  for a,b in [((745,60),(470,215)),((745,60),(1040,215)),((470,215),(470,370))]:edge(a,b)
  for x,y,t in [(745,60,'1  权值5'),(470,215,'2  权值4'),(1040,215,'3  权值3'),(470,370,'4  权值6')]:node(x,y,t)
  d.text((1040,350),'相连的两点不能同时选择',font=fsmall,anchor='mm',fill='black')
 im.save(path);return path

for k in ['recursion','dag','tree']:diagram(k)

def paragraph(doc,text,style=None,bold=None):
 p=doc.add_paragraph(style=style);r=p.add_run(text)
 # None inherits the paragraph style. False would cancel bold headings.
 if bold is not None:r.bold=bold
 p.paragraph_format.widow_control=True
 return p

def table(doc,rows):
 n=len(rows[0]); t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 total=6.95
 # compact comparison matrices rather than narrative containers
 if n==2: widths=[2.25,total-2.25]
 elif n==3: widths=[1.4,2.75,2.8]
 elif n>=5: widths=[1.3]+[(total-1.3)/(n-1)]*(n-1)
 else:widths=[total/n]*n
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders)
 for ri,data in enumerate(rows):
  cells=t.add_row().cells
  for ci,txt in enumerate(data):
   c=cells[ci];c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.style=doc.styles['Table Text']
   # Keep compact teaching tables intact; longer tables repeat their header.
   p.paragraph_format.keep_with_next=len(rows)<=5 and ri<len(rows)-1
   p.add_run(txt).bold=ri==0
   tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   if ri==0:
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EDF2');tcpr.append(fill)
  trpr=t.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if ri==0:
   hdr=OxmlElement('w:tblHeader');trpr.append(hdr)
 paragraph(doc,'','Tiny')
 return t

for vol,cn in enumerate(['一','二','三','四']):
 lines=(ROOT/f'内容{cn}.txt').read_text().splitlines()
 doc=Document();sec=doc.sections[0]
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.78);sec.right_margin=Inches(.77)
 sec.header_distance=Inches(.25);sec.footer_distance=Inches(.27)
 sty=doc.styles
 set_font(sty['Normal'],LATIN,11.5);sty['Normal'].paragraph_format.line_spacing=Pt(18.5);sty['Normal'].paragraph_format.space_after=Pt(6)
 for key,size,leading,before,after in [('Title',27,38,0,14),('Subtitle',16,26,0,14),('Heading 1',21,30,24,14),('Heading 2',14.5,23,18,8),('Heading 3',12,20,12,6)]:
  set_font(sty[key],LATIN,size,True,east_asia=HEAD_FONT)
  sty[key].paragraph_format.line_spacing=Pt(leading)
  sty[key].paragraph_format.space_before=Pt(before);sty[key].paragraph_format.space_after=Pt(after)
  sty[key].paragraph_format.keep_with_next=True
 for key in ['Title','Subtitle']:
  sty[key].paragraph_format.space_before=Pt(0)
  sty[key].paragraph_format.keep_with_next=True
 for name,font,size,lead in [('Code',MONO,10.5,1.16),('Formula',MONO,11.0,1.35),('Table Text',LATIN,10.5,1.2),('Tiny',LATIN,3,1),('TOC Text',LATIN,11,1.2)]:
  if name not in sty:sty.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
  set_font(sty[name],font,size)
  sty[name].paragraph_format.line_spacing=Pt(16 if name in ['Code','Table Text','TOC Text'] else 18);sty[name].paragraph_format.space_after=Pt(5)
 sty['Tiny'].paragraph_format.line_spacing=Pt(3)
 sty['Code'].paragraph_format.left_indent=Inches(.13)
 sty['Formula'].paragraph_format.left_indent=Inches(.1)
 sty['Formula'].paragraph_format.space_before=Pt(2)
 header=sec.header.paragraphs[0];header.text=f'动态规划学习材料  第{cn}册';header.style=sty['Normal']
 header.runs[0].font.size=Pt(8);header.paragraph_format.space_after=Pt(0)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=footer.add_run('2026年10月12日    ');r.font.size=Pt(8)
 add_field(footer,'PAGE');footer.paragraph_format.space_after=Pt(0)
 doc.core_properties.title=lines[0].split(' ',1)[1]+' '+lines[1].split(' ',1)[1]
 doc.core_properties.subject='从零起步到提高组的动态规划学习与验收'
 doc.core_properties.author='家庭学习教练';doc.core_properties.keywords='动态规划 DP GESP CSP-S'
 firsth=next(i for i,x in enumerate(lines) if x.startswith('@h1 '))
 for line in lines[:firsth]:
  if line.startswith('@title '):paragraph(doc,line[7:],'Title')
  elif line.startswith('@subtitle '):paragraph(doc,line[10:],'Subtitle')
  elif line.startswith('@intro '):paragraph(doc,line[7:],bold=True)
  elif line.startswith('@h2 '):paragraph(doc,line[4:],'Heading 2')
  elif line:paragraph(doc,line)
 # TOC always begins next page; links usable in Word and PDF
 doc.add_page_break();paragraph(doc,'本册目录','Heading 1')
 headings=[x[4:] for x in lines if x.startswith('@h1 ')]
 for idx,h in enumerate(headings):
  p=doc.add_paragraph(style='TOC Text');p.paragraph_format.tab_stops.add_tab_stop(Inches(6.55),WD_ALIGN_PARAGRAPH.RIGHT)
  link(p,h,f'chapter_{idx}')
  page=PAGE_MAP.get(NAMES[vol],{}).get(h,'')
  if page:p.add_run('\t'+str(page))
 paragraph(doc,'目录可点击跳转。打印阅读时可按章选择页码。','Table Text')
 if vol==3:paragraph(doc,'本册含核验结论，请在孩子尝试后查看相应条目。','Table Text')
 idx=firsth;chapter=0;lesson_number=None;subsection=0
 while idx<len(lines):
  line=lines[idx];idx+=1
  if not line:continue
  if line.startswith('@h1 '):
   p=paragraph(doc,line[4:],'Heading 1');p.paragraph_format.page_break_before=(chapter==0 or vol==2)
   add_bookmark(p,f'chapter_{chapter}',chapter+1);chapter+=1
   match=re.match(r'第(\d+)章',line[4:]);lesson_number=int(match[1]) if match and vol<2 else None;subsection=0
  elif line.startswith('@h2 '):
   subsection+=1
   label=(f'{lesson_number}.{subsection}  ' if lesson_number is not None else '')+line[4:]
   paragraph(doc,label,'Heading 2')
  elif line.startswith('@table '):
   rows=[line[7:].split('|')]
   while lines[idx]!='@endtable':rows.append(lines[idx].split('|'));idx+=1
   idx+=1
   assert all(len(x)==len(rows[0]) for x in rows),rows
   table(doc,rows)
  elif line.startswith('@formula '):
   # Array-style program relations deliberately use editable linear notation.
   paragraph(doc,line[9:],'Formula')
  elif line.startswith('@code ') or line=='@sample':
   if line.startswith('@code '):paragraph(doc,line[6:],'Heading 3')
   buf=[]
   while lines[idx]!='@endcode':buf.append(lines[idx]);idx+=1
   idx+=1
   for j,t in enumerate(buf):
    p=paragraph(doc,t,'Code');p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.keep_with_next=j<len(buf)-1
  elif line.startswith('@diagram '):
   k=line.split()[1];p=doc.add_paragraph();p.paragraph_format.line_spacing=1.0;p.add_run().add_picture(str(TMP/(k+'.png')),width=Inches(6.75));p.paragraph_format.keep_with_next=True
   # Give illustration useful alternative text.
   descr={'recursion':'递归调用中的重复子问题','dag':'两条从点1到点4的有向路径','tree':'四个节点的树与快乐值'}[k]
   for e in p._p.xpath('.//wp:docPr'):e.set('descr',descr)
  elif line.startswith('@task '):
   paragraph(doc,'独立检查','Heading 2')
   paragraph(doc,line[6:])
   for label in ['我的状态和来源','我设计的关键测试','第一次偏差与修改依据']:
    paragraph(doc,label+'  __________________________________________________','Table Text')
  elif line.startswith('@url '):
   p=paragraph(doc,line[5:],'Table Text');p.paragraph_format.space_after=Pt(4)
  elif line.startswith('@'):raise ValueError(line)
  else:paragraph(doc,line)
 # Remove inherited Word title rules and decorative borders.
 for root in [doc._element,doc.styles.element]:
  for border in root.xpath('.//w:pBdr'):
   border.getparent().remove(border)
 path=ROOT/(NAMES[vol]+'.docx');doc.save(path);print(path)
print('四册可编辑版生成完成')
